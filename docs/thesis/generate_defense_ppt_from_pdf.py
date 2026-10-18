#!/usr/bin/env python3
"""Generate editable defense PPT matching 点云数据滤波程序设计答辩PPT.pdf."""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

PRIMARY = RGBColor(0x1A, 0x56, 0x8E)
ACCENT = RGBColor(0x2E, 0x86, 0xC1)
TEXT = RGBColor(0x2C, 0x3E, 0x50)
MUTED = RGBColor(0x66, 0x66, 0x66)
LIGHT = RGBColor(0xF4, 0xF8, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x22, 0xB1, 0x7B)
GREY = RGBColor(0x7E, 0x8B, 0x9A)

SCRIPT_DIR = Path(__file__).resolve().parent
PDF_SLIDES = SCRIPT_DIR / "pdf_slides"
ASSET_DIR = SCRIPT_DIR / "ppt_assets"
OUTPUT_PATH = SCRIPT_DIR / "点云数据滤波程序设计答辩PPT.pptx"


def set_slide_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_top_bar(slide) -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = PRIMARY
    bar.line.fill.background()


def add_title_bar(slide, title: str) -> None:
    add_top_bar(slide)
    title_box = slide.shapes.add_textbox(Inches(0.55), Inches(0.25), Inches(8.9), Inches(0.55))
    title_frame = title_box.text_frame
    title_frame.text = title
    title_paragraph = title_frame.paragraphs[0]
    title_paragraph.font.size = Pt(28)
    title_paragraph.font.bold = True
    title_paragraph.font.color.rgb = PRIMARY
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.85), Inches(8.9), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()


def add_bullets(
    slide,
    items: list[str],
    left: float = 0.55,
    top: float = 1.15,
    width: float = 8.9,
    height: float = 5.8,
    font_size: int = 18,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    text_frame = box.text_frame
    text_frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = item
        paragraph.font.size = Pt(font_size)
        paragraph.font.color.rgb = TEXT
        paragraph.space_after = Pt(10)


def add_card(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    title: str,
    body: str | list[str],
    title_size: int = 18,
    body_size: int = 14,
) -> None:
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = RGBColor(0xD6, 0xE4, 0xF0)
    card.shadow.inherit = False

    box = slide.shapes.add_textbox(Inches(left + 0.15), Inches(top + 0.12), Inches(width - 0.3), Inches(height - 0.2))
    text_frame = box.text_frame
    text_frame.word_wrap = True
    title_paragraph = text_frame.paragraphs[0]
    title_paragraph.text = title
    title_paragraph.font.size = Pt(title_size)
    title_paragraph.font.bold = True
    title_paragraph.font.color.rgb = PRIMARY
    title_paragraph.space_after = Pt(8)

    body_items = body if isinstance(body, list) else [body]
    for body_item in body_items:
        paragraph = text_frame.add_paragraph()
        paragraph.text = body_item
        paragraph.font.size = Pt(body_size)
        paragraph.font.color.rgb = TEXT
        paragraph.space_after = Pt(6)


def add_progress_bar(slide, left: float, top: float, width: float, label: str, percent: int, color: RGBColor) -> None:
    label_box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(2.8), Inches(0.35))
    label_paragraph = label_box.text_frame.paragraphs[0]
    label_paragraph.text = label
    label_paragraph.font.size = Pt(14)
    label_paragraph.font.color.rgb = TEXT

    track = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left + 2.9),
        Inches(top + 0.02),
        Inches(width),
        Inches(0.32),
    )
    track.fill.solid()
    track.fill.fore_color.rgb = RGBColor(0xE8, 0xED, 0xF2)
    track.line.fill.background()

    fill_width = max(width * percent / 100, 0.8)
    fill = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left + 2.9),
        Inches(top + 0.02),
        Inches(fill_width),
        Inches(0.32),
    )
    fill.fill.solid()
    fill.fill.fore_color.rgb = color
    fill.line.fill.background()

    value_box = slide.shapes.add_textbox(Inches(left + 2.95), Inches(top + 0.02), Inches(fill_width - 0.1), Inches(0.32))
    value_frame = value_box.text_frame
    value_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    value_paragraph = value_frame.paragraphs[0]
    value_paragraph.text = f"{label.split()[0]} 效率 ({percent}%)"
    value_paragraph.font.size = Pt(12)
    value_paragraph.font.bold = True
    value_paragraph.font.color.rgb = WHITE


def crop_pdf_slide(slide_number: int, box: tuple[int, int, int, int], output_name: str) -> Path:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    source_path = PDF_SLIDES / f"slide_{slide_number:02d}.png"
    output_path = ASSET_DIR / output_name
    if output_path.exists():
        return output_path
    image = Image.open(source_path)
    cropped = image.crop(box)
    cropped.save(output_path)
    return output_path


def add_picture(slide, image_path: Path, left: float, top: float, width: float, height: float) -> None:
    if image_path.exists():
        slide.shapes.add_picture(str(image_path), Inches(left), Inches(top), width=Inches(width), height=Inches(height))


def build_presentation() -> Presentation:
    presentation = Presentation()
    presentation.slide_width = Inches(10)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    # Slide 1 — Cover
    slide = presentation.slides.add_slide(blank)
    set_slide_bg(slide, LIGHT)
    add_top_bar(slide)
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(1.0))
    title_paragraph = title_box.text_frame.paragraphs[0]
    title_paragraph.text = "点云数据滤波程序设计"
    title_paragraph.font.size = Pt(40)
    title_paragraph.font.bold = True
    title_paragraph.font.color.rgb = PRIMARY
    title_paragraph.alignment = PP_ALIGN.CENTER

    sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.55), Inches(8.4), Inches(0.6))
    sub_paragraph = sub_box.text_frame.paragraphs[0]
    sub_paragraph.text = "江西理工大学 · 测绘工程 · 本科毕业设计答辩"
    sub_paragraph.font.size = Pt(20)
    sub_paragraph.font.color.rgb = ACCENT
    sub_paragraph.alignment = PP_ALIGN.CENTER

    divider = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.2), Inches(3.35), Inches(1.6), Inches(0.04))
    divider.fill.solid()
    divider.fill.fore_color.rgb = ACCENT
    divider.line.fill.background()

    info_box = slide.shapes.add_textbox(Inches(0.8), Inches(3.75), Inches(8.4), Inches(1.8))
    info_frame = info_box.text_frame
    info_lines = [
        "答辩人：张钧炀（学号：1220220997）",
        "指导老师：敖建锋 副教授",
        "答辩时间：2026年5月",
    ]
    for index, line in enumerate(info_lines):
        paragraph = info_frame.paragraphs[0] if index == 0 else info_frame.add_paragraph()
        paragraph.text = line
        paragraph.font.size = Pt(18)
        paragraph.font.color.rgb = TEXT
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.space_after = Pt(8)

    # Slide 2 — TOC title
    slide = presentation.slides.add_slide(blank)
    set_slide_bg(slide, LIGHT)
    add_top_bar(slide)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.35), Inches(2.5), Inches(1.3), Inches(0.08))
    accent.fill.solid()
    accent.fill.fore_color.rgb = ACCENT
    accent.line.fill.background()
    toc_title = slide.shapes.add_textbox(Inches(1.0), Inches(2.75), Inches(8.0), Inches(0.8))
    toc_title.text_frame.text = "汇报目录"
    toc_title.text_frame.paragraphs[0].font.size = Pt(36)
    toc_title.text_frame.paragraphs[0].font.bold = True
    toc_title.text_frame.paragraphs[0].font.color.rgb = PRIMARY
    toc_title.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    toc_sub = slide.shapes.add_textbox(Inches(1.0), Inches(3.55), Inches(8.0), Inches(0.5))
    toc_sub.text_frame.text = "TABLE OF CONTENTS"
    toc_sub.text_frame.paragraphs[0].font.size = Pt(16)
    toc_sub.text_frame.paragraphs[0].font.color.rgb = MUTED
    toc_sub.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    # Slide 3 — TOC items
    slide = presentation.slides.add_slide(blank)
    set_slide_bg(slide, LIGHT)
    add_title_bar(slide, "汇报目录")
    toc_items = [
        "第一部分：研究背景与意义 — 探讨点云噪声来源及滤波的必要性",
        "第二部分：研究问题与目标 — 提出针对测绘工程的预处理方案",
        "第三部分：核心研究方法 — 详述空间索引、SOR 与 ROR 滤波机制",
        "第四部分：实验结果分析 — 对比分析算法去除率及结构保真度",
        "第五部分：创新总结与展望 — 工程应用价值与未来演进方向",
    ]
    add_bullets(slide, toc_items, top=1.2, font_size=20)

    # Slide 4 — Background
    slide = presentation.slides.add_slide(blank)
    set_slide_bg(slide, LIGHT)
    add_title_bar(slide, "研究背景与意义")
    background_image = crop_pdf_slide(4, (980, 260, 1880, 980), "background_pointcloud.png")
    add_bullets(
        slide,
        [
            "随着激光雷达（LiDAR）和三维扫描技术的发展，点云数据已成为地形测绘、城市三维建模等领域的核心数据源。",
            "在实际采集过程中，受仪器精度、环境遮挡及物体反射率偏差等影响，原始点云中往往含有大量孤立点和噪声点。",
            "无效的噪声数据会降低点云质量，严重影响后续的点云配准、地物分类以及曲面重构等关键步骤的准确性。",
            "传统滤波方法面临参数敏感、缺乏自动化以及无法高效应对海量点云数据等工程痛点。",
        ],
        width=5.2,
        height=5.5,
        font_size=16,
    )
    add_picture(slide, background_image, 5.95, 1.15, 3.5, 5.2)

    # Slide 5 — Goals
    slide = presentation.slides.add_slide(blank)
    set_slide_bg(slide, LIGHT)
    add_title_bar(slide, "研究目标与核心问题")
    cards = [
        (
            0.45,
            "核心目标",
            "从测绘工程应用角度出发，开发一套自动化、参数化的点云滤波预处理程序，解决现有脚本操作不便的问题。",
        ),
        (
            3.45,
            "算法落地",
            "探索并编程实现统计离群滤波（SOR）与半径离群滤波（ROR），通过对比测试寻找最优的噪声剔除方案。",
        ),
        (
            6.45,
            "海量数据应对",
            "引入空间索引树提升查找效率，并通过带有重叠区域的 XY 平面分块策略，解决大文件处理时的内存瓶颈。",
        ),
    ]
    for left, title, body in cards:
        add_card(slide, left, 1.25, 2.85, 5.4, title, body)

    # Slide 6 — Spatial index
    slide = presentation.slides.add_slide(blank)
    set_slide_bg(slide, LIGHT)
    add_title_bar(slide, "核心机制：空间索引与邻域搜索")
    add_card(
        slide,
        0.45,
        1.25,
        4.35,
        5.5,
        "KD-Tree（K 维树）",
        [
            "基于数据维度方差划分的二叉搜索树结构，可快速排除空间中较远的无关区域。",
            "• 按 X、Y、Z 坐标维度依次构建树状索引",
            "• 在低维（3D 坐标）中极大加速 K 近邻（K-NN）搜索",
            "• 本程序中主要用于加速 SOR 统计离群点滤波的距离计算",
        ],
    )
    add_card(
        slide,
        5.15,
        1.25,
        4.35,
        5.5,
        "Ball Tree（球树）",
        [
            "使用多维超球体划分空间结构，查询时基于中心距离与半径的距离下界进行节点剪枝。",
            "• 克服 KD-Tree 在高维或复杂密度分布下的性能退化",
            "• 对基于固定半径的邻域查询具有效率优势",
            "• 本程序中重点辅助 ROR 半径离群点滤波及密度聚类算法",
        ],
    )

    # Slide 7 — SOR
    slide = presentation.slides.add_slide(blank)
    set_slide_bg(slide, LIGHT)
    add_title_bar(slide, "算法一：统计离群点滤波（SOR）")
    add_bullets(
        slide,
        [
            "基本原理：假设正常点云的局部距离服从高斯分布，根据各点与其 K 个近邻点的平均距离进行统计判定。",
            "公式：dᵢ = (1/k) Σⱼ₌₁ᵏ ‖pᵢ − pᵢⱼ‖₂",
            "判断准则：计算全体数据的均值 μ 和标准差 σ。若 dᵢ > μ + ασ，则判定并剔除该离群点。",
            "算法特点：对全局范围内的随机噪声具有强大的鲁棒性，能极好地保持整体结构连贯性。",
        ],
        width=5.0,
        height=5.5,
        font_size=15,
    )
    sor_image = crop_pdf_slide(7, (980, 250, 1880, 980), "sor_compare.png")
    add_picture(slide, sor_image, 5.75, 1.15, 3.7, 5.2)

    # Slide 8 — ROR
    slide = presentation.slides.add_slide(blank)
    set_slide_bg(slide, LIGHT)
    add_title_bar(slide, "算法二：半径离群点滤波（ROR）")
    add_bullets(
        slide,
        [
            "基于密度的局部过滤",
            "基本思想：不同于 SOR 依赖平均距离，ROR 完全基于点云的局部空间密度特征来识别异常点，专门用于剔除孤立的散乱噪声。",
            "执行流程：",
            "• 以空间点 pᵢ 为球心，设定固定搜索半径 r",
            "• 统计该超球体内包含的邻域点数量 |N(pᵢ)|",
            "• 若 |N(pᵢ)| < T，则该点被认定为缺乏有效支撑的离群噪声点",
            "应用场景：对局部稀疏噪声极为敏感，非常适合处理边缘悬浮点与孤立噪声。",
        ],
        width=5.0,
        height=5.5,
        font_size=15,
    )
    ror_image = crop_pdf_slide(8, (980, 250, 1880, 980), "ror_compare.png")
    add_picture(slide, ror_image, 5.75, 1.15, 3.7, 5.2)

    # Slide 9 — Extensions
    slide = presentation.slides.add_slide(blank)
    set_slide_bg(slide, LIGHT)
    add_title_bar(slide, "拓展与高级处理模块")
    extension_cards = [
        (
            0.35,
            "点云分割技术",
            "引入密度聚类（DBSCAN）和区域生长算法，将滤波后的纯净点云按几何或密度特征划分为独立的目标区域，为后续提取提供支撑。",
        ),
        (
            3.45,
            "语义分类与裁切",
            "提取法线、曲率等几何特征后进行分类标识，并实现基于语义标签的点云子集裁切，高度契合测绘地物提取工作流。",
        ),
        (
            6.55,
            "海量数据重叠分块",
            "针对超大体积 LAS 文件，采取带重叠缓冲区的 XY 平面分块（Tiled）处理策略，有效缓解内存溢出并消除边界滤波假影。",
        ),
    ]
    extension_images = [
        crop_pdf_slide(9, (70, 250, 620, 620), "segmentation.png"),
        crop_pdf_slide(9, (680, 250, 1230, 620), "classification_legend.png"),
        crop_pdf_slide(9, (1290, 250, 1840, 620), "tiling_scene.png"),
    ]
    for index, ((left, title, body), image_path) in enumerate(zip(extension_cards, extension_images)):
        add_picture(slide, image_path, left, 1.15, 2.85, 2.0)
        add_card(slide, left, 3.25, 2.85, 3.5, title, body, body_size=13)

    # Slide 10 — Results
    slide = presentation.slides.add_slide(blank)
    set_slide_bg(slide, LIGHT)
    add_title_bar(slide, "结果分析：综合滤波性能对比")
    add_progress_bar(slide, 0.55, 1.45, 6.0, "SOR", 92, PRIMARY)
    add_progress_bar(slide, 0.55, 2.05, 6.0, "ROR", 65, GREY)
    label_box = slide.shapes.add_textbox(Inches(0.55), Inches(1.15), Inches(8.5), Inches(0.3))
    label_box.text_frame.text = "处理全局分布随机噪声"
    label_box.text_frame.paragraphs[0].font.size = Pt(15)
    label_box.text_frame.paragraphs[0].font.bold = True

    add_progress_bar(slide, 0.55, 3.05, 6.0, "SOR", 75, ACCENT)
    add_progress_bar(slide, 0.55, 3.65, 6.0, "ROR", 95, GREEN)
    label_box2 = slide.shapes.add_textbox(Inches(0.55), Inches(2.75), Inches(8.5), Inches(0.3))
    label_box2.text_frame.text = "处理局部稀疏孤立噪声"
    label_box2.text_frame.paragraphs[0].font.size = Pt(15)
    label_box2.text_frame.paragraphs[0].font.bold = True

    summary_box = slide.shapes.add_textbox(Inches(0.55), Inches(4.55), Inches(8.9), Inches(1.8))
    summary_paragraph = summary_box.text_frame.paragraphs[0]
    summary_paragraph.text = (
        "实验数据显示：SOR 算法在应对全局性、大范围离群噪声时具有卓越的鲁棒性且能更好保护地形结构；"
        "而 ROR 算法则对于局部游离的稀疏噪声表现出更高的敏感度和极强的清除能力。二者相辅相成。"
    )
    summary_paragraph.font.size = Pt(15)
    summary_paragraph.font.italic = True
    summary_paragraph.font.color.rgb = MUTED

    # Slide 11 — Innovation
    slide = presentation.slides.add_slide(blank)
    set_slide_bg(slide, LIGHT)
    add_title_bar(slide, "工程创新点总结")
    innovations = [
        (
            "无缝的精度保留机制",
            "在 LAS 文件读写处理中，严格保持原始头信息中的比例尺（Scale）和偏移量（Offset）不变，"
            "确保滤波前后坐标精度零损失，完美兼容现存测绘商业软件。",
        ),
        (
            "轻量高效的技术栈设计",
            "摒弃配置繁杂的传统 C++ 桌面库体系，采用 laspy + NumPy + SciPy + scikit-learn 组合，"
            "在保证计算性能的同时极大降低了工程部署门槛。",
        ),
        (
            "重叠分块内存优化策略",
            "面对 GB 级别超大点云文件，创新实现带有冗余缓冲区的 XY 分块并行处理，"
            "根治内存峰值崩溃问题，并运用缓冲区重叠消除分界线的接缝瑕疵。",
        ),
        (
            "极具实战价值的工作流",
            "程序集成滤波、降噪、分割与分类提取全套预处理流程，"
            "可直接作为地面提取、数字孪生及分类建模等高阶任务的上游标准预处理管线。",
        ),
    ]
    top_positions = [1.15, 2.55, 3.95, 5.35]
    for (title, body), top in zip(innovations, top_positions):
        add_card(slide, 0.55, top, 8.9, 1.15, title, body, title_size=17, body_size=13)

    # Slide 12 — Conclusion
    slide = presentation.slides.add_slide(blank)
    set_slide_bg(slide, LIGHT)
    add_title_bar(slide, "研究结论与未来展望")
    conclusion_image = crop_pdf_slide(12, (70, 250, 620, 980), "conclusion_pointcloud.png")
    add_picture(slide, conclusion_image, 0.55, 1.2, 3.2, 5.3)
    add_card(
        slide,
        4.0,
        1.2,
        5.45,
        2.4,
        "主要研究结论",
        "本课题成功研发了一款面向实际测绘工程应用的点云滤波预处理程序。实验证明，该程序能够高效且智能地"
        "剔除不同特征的异常噪声点，同时极其出色地保持了原始点云的地物细节与结构连贯性，具有较高的应用转化价值。",
        body_size=15,
    )
    add_card(
        slide,
        4.0,
        3.85,
        5.45,
        2.65,
        "未来发展展望",
        "随着 AI 的演进，未来计划在管线中引入 PointNet 等深度学习模型，以实现动态复杂场景下的语义级智能滤波。"
        "同时，将探索基于 GPU（CUDA）的并行计算底层架构重构，进一步提升海量城市场景点云的实时处理速率。",
        body_size=15,
    )

    # Slide 13 — Thanks
    slide = presentation.slides.add_slide(blank)
    set_slide_bg(slide, LIGHT)
    add_top_bar(slide)
    thanks_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(8.0), Inches(1.0))
    thanks_paragraph = thanks_box.text_frame.paragraphs[0]
    thanks_paragraph.text = "感谢聆听，敬请批评指正！"
    thanks_paragraph.font.size = Pt(36)
    thanks_paragraph.font.bold = True
    thanks_paragraph.font.color.rgb = PRIMARY
    thanks_paragraph.alignment = PP_ALIGN.CENTER
    qa_box = slide.shapes.add_textbox(Inches(1.0), Inches(3.2), Inches(8.0), Inches(0.5))
    qa_paragraph = qa_box.text_frame.paragraphs[0]
    qa_paragraph.text = "Q & A Session"
    qa_paragraph.font.size = Pt(20)
    qa_paragraph.font.color.rgb = ACCENT
    qa_paragraph.alignment = PP_ALIGN.CENTER
    footer_box = slide.shapes.add_textbox(Inches(1.0), Inches(5.0), Inches(8.0), Inches(1.2))
    footer_frame = footer_box.text_frame
    footer_lines = [
        "汇报人：张钧炀  |  指导教师：敖建锋 副教授",
        "江西理工大学 · 土木与测绘工程学院",
    ]
    for index, line in enumerate(footer_lines):
        paragraph = footer_frame.paragraphs[0] if index == 0 else footer_frame.add_paragraph()
        paragraph.text = line
        paragraph.font.size = Pt(16)
        paragraph.font.color.rgb = MUTED
        paragraph.alignment = PP_ALIGN.CENTER

    # Slide 14 — Image sources 1
    slide = presentation.slides.add_slide(blank)
    set_slide_bg(slide, LIGHT)
    add_title_bar(slide, "Image Sources")
    sources_page_one = [
        ("https://images.ctfassets.net/go54bjdzbrgi/fNuDew3tlWg2OOywYOIKM/ebef71a9e739dd33efc5297dce42034c/Pix4D_construction_surveys_BIM000b.jpg", "www.pix4d.com"),
        ("https://encyclopedia.pub/media/item_content/202305/646d59f4dc53aelectronics-12-02150-g001.png", "encyclopedia.pub"),
        ("https://pub.mdpi-res.com/data/data-08-00149/article_deploy/html/images/data-08-00149-g003.png?1696669431", "www.mdpi.com"),
        ("https://miro.medium.com/1*BXhxw5NzWR40y-7bPupwFQ.gif", "medium.com"),
        ("https://s3.amazonaws.com/cdn.freshdesk.com/data/helpdesk/attachments/production/31031488770/original/X0DU_HOQZQaJZM4IKhMq8fayBmTfvmULkw.png?1623162812", "agisoft.freshdesk.com"),
        ("https://segments.ai/wp-content/uploads/2024/02/Mastering-point-clouds-data.jpg", "segments.ai"),
    ]
    source_images = [
        crop_pdf_slide(14, (70, 180, 250, 320), "source_1.png"),
        crop_pdf_slide(14, (70, 360, 250, 500), "source_2.png"),
        crop_pdf_slide(14, (70, 540, 250, 680), "source_3.png"),
        crop_pdf_slide(14, (70, 720, 250, 860), "source_4.png"),
        crop_pdf_slide(14, (70, 900, 250, 1040), "source_5.png"),
        crop_pdf_slide(14, (70, 1080, 250, 1220), "source_6.png"),
    ]
    for index, ((url, site), image_path) in enumerate(zip(sources_page_one, source_images)):
        top = 1.05 + index * 0.95
        add_picture(slide, image_path, 0.55, top, 0.85, 0.65)
        text_box = slide.shapes.add_textbox(Inches(1.55), Inches(top), Inches(7.9), Inches(0.65))
        text_frame = text_box.text_frame
        text_frame.word_wrap = True
        url_paragraph = text_frame.paragraphs[0]
        url_paragraph.text = url
        url_paragraph.font.size = Pt(10)
        url_paragraph.font.color.rgb = TEXT
        site_paragraph = text_frame.add_paragraph()
        site_paragraph.text = f"Source: {site}"
        site_paragraph.font.size = Pt(10)
        site_paragraph.font.color.rgb = ACCENT

    # Slide 15 — Image sources 2
    slide = presentation.slides.add_slide(blank)
    set_slide_bg(slide, LIGHT)
    add_title_bar(slide, "Image Sources")
    source_image = crop_pdf_slide(15, (70, 320, 250, 460), "source_7.png")
    add_picture(slide, source_image, 0.55, 1.5, 0.85, 0.65)
    text_box = slide.shapes.add_textbox(Inches(1.55), Inches(1.5), Inches(7.9), Inches(0.8))
    text_frame = text_box.text_frame
    text_frame.paragraphs[0].text = "https://aecmag.com/wp-content/uploads/2023/03/beforeclean1500-1024x587.jpg"
    text_frame.paragraphs[0].font.size = Pt(11)
    text_frame.paragraphs[0].font.color.rgb = TEXT
    source_paragraph = text_frame.add_paragraph()
    source_paragraph.text = "Source: aecmag.com"
    source_paragraph.font.size = Pt(11)
    source_paragraph.font.color.rgb = ACCENT

    return presentation


def main() -> None:
    presentation = build_presentation()
    presentation.save(OUTPUT_PATH)
    print(f"Generated: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
