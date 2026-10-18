#!/usr/bin/env python3
"""Clean defense PPT for 张钧炀 — surveying-focused, ~10 min, consistent layout."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUTPUT_PATH = Path("/Users/kyle/Desktop/毕业设计答辩PPT-测绘223张钧炀.pptx")
REPO_IMAGES = Path(__file__).resolve().parent / "images"
DOCX_IMAGES = Path(__file__).resolve().parent / "extracted_docx_media"

PRIMARY = RGBColor(0x3D, 0x6B, 0x8A)
PRIMARY_DARK = RGBColor(0x2C, 0x4F, 0x66)
ACCENT = RGBColor(0xB8, 0x5A, 0x4A)
TEXT = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xF7, 0xF9, 0xFB)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN_L = Inches(0.55)
CONTENT_TOP = Inches(1.15)
TITLE_H = Inches(0.85)


def img(name: str) -> Path:
    candidate = DOCX_IMAGES / name
    if candidate.exists():
        return candidate
    return REPO_IMAGES / name


def set_bg(slide, color: RGBColor = BG) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_header_bar(slide, title: str) -> None:
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), SLIDE_W, TITLE_H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = PRIMARY
    bar.line.fill.background()
    text_frame = bar.text_frame
    text_frame.text = title
    paragraph = text_frame.paragraphs[0]
    paragraph.font.size = Pt(26)
    paragraph.font.bold = True
    paragraph.font.color.rgb = WHITE
    paragraph.alignment = PP_ALIGN.LEFT
    text_frame.margin_left = MARGIN_L
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE


def add_footer(slide, text: str) -> None:
    box = slide.shapes.add_textbox(MARGIN_L, Inches(7.15), Inches(12), Inches(0.25))
    frame = box.text_frame
    frame.text = text
    paragraph = frame.paragraphs[0]
    paragraph.font.size = Pt(10)
    paragraph.font.color.rgb = MUTED


def add_bullets(
    slide,
    items: list[str],
    left=Inches(0.65),
    top=CONTENT_TOP,
    width=Inches(5.8),
    height=Inches(5.6),
    font_size: int = 20,
) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.size = Pt(font_size)
        paragraph.font.color.rgb = TEXT
        paragraph.space_after = Pt(10)


def add_picture(
    slide,
    image_path: Path,
    left,
    top,
    width,
    height,
    caption: str | None = None,
) -> None:
    if image_path.exists():
        slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)
    else:
        placeholder = slide.shapes.add_shape(1, left, top, width, height)
        placeholder.fill.solid()
        placeholder.fill.fore_color.rgb = RGBColor(0xE8, 0xED, 0xF2)
        placeholder.line.color.rgb = PRIMARY
        text_frame = placeholder.text_frame
        text_frame.text = f"【待补图】\n{caption or image_path.name}"
        text_frame.paragraphs[0].font.size = Pt(14)
        text_frame.paragraphs[0].font.color.rgb = ACCENT
        text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    if caption:
        cap = slide.shapes.add_textbox(left, top + height + Inches(0.05), width, Inches(0.3))
        cap_frame = cap.text_frame
        cap_frame.text = caption
        cap_paragraph = cap_frame.paragraphs[0]
        cap_paragraph.font.size = Pt(11)
        cap_paragraph.font.color.rgb = MUTED
        cap_paragraph.alignment = PP_ALIGN.CENTER


def slide_cover(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, PRIMARY_DARK)
    top_band = slide.shapes.add_shape(1, Inches(0), Inches(2.2), SLIDE_W, Inches(2.8))
    top_band.fill.solid()
    top_band.fill.fore_color.rgb = PRIMARY
    top_band.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(11.5), Inches(1.2))
    title_frame = title_box.text_frame
    title_frame.text = "点云数据滤波程序设计"
    title_paragraph = title_frame.paragraphs[0]
    title_paragraph.font.size = Pt(40)
    title_paragraph.font.bold = True
    title_paragraph.font.color.rgb = WHITE
    title_paragraph.alignment = PP_ALIGN.CENTER

    sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(3.7), Inches(11.5), Inches(0.6))
    sub_frame = sub_box.text_frame
    sub_frame.text = "本科毕业设计答辩"
    sub_paragraph = sub_frame.paragraphs[0]
    sub_paragraph.font.size = Pt(22)
    sub_paragraph.font.color.rgb = RGBColor(0xD6, 0xE8, 0xF3)
    sub_paragraph.alignment = PP_ALIGN.CENTER

    info_box = slide.shapes.add_textbox(Inches(0.8), Inches(5.0), Inches(11.5), Inches(1.6))
    info_frame = info_box.text_frame
    info_frame.text = (
        "低空技术学院  测绘工程223班\n"
        "答辩人：张钧炀          导师：敖建锋"
    )
    for paragraph in info_frame.paragraphs:
        paragraph.font.size = Pt(20)
        paragraph.font.color.rgb = WHITE
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.space_after = Pt(8)


def slide_contents(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header_bar(slide, "答辩内容")
    items = [
        "1. 研究背景与意义",
        "2. 点云滤波理论基础（SOR / ROR）",
        "3. 程序设计与实现思路",
        "4. 实验结果与分析",
        "5. 结论与展望",
    ]
    add_bullets(slide, items, top=Inches(1.8), font_size=24)
    add_footer(slide, "测绘工程223班  张钧炀")


def slide_section(prs: Presentation, number: str, title: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, PRIMARY)
    number_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(3), Inches(1.5))
    number_frame = number_box.text_frame
    number_frame.text = number
    number_paragraph = number_frame.paragraphs[0]
    number_paragraph.font.size = Pt(72)
    number_paragraph.font.bold = True
    number_paragraph.font.color.rgb = RGBColor(0x8A, 0xAD, 0xC4)

    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(3.5), Inches(11), Inches(1.0))
    title_frame = title_box.text_frame
    title_frame.text = title
    title_paragraph = title_frame.paragraphs[0]
    title_paragraph.font.size = Pt(36)
    title_paragraph.font.bold = True
    title_paragraph.font.color.rgb = WHITE


def slide_text_image(
    prs: Presentation,
    title: str,
    bullets: list[str],
    image_path: Path,
    caption: str,
    footer: str = "",
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header_bar(slide, title)
    add_bullets(slide, bullets, width=Inches(5.5))
    add_picture(
        slide,
        image_path,
        Inches(6.5),
        CONTENT_TOP,
        Inches(6.2),
        Inches(5.3),
        caption,
    )
    if footer:
        add_footer(slide, footer)


def slide_full_image(prs: Presentation, title: str, image_path: Path, caption: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header_bar(slide, title)
    add_picture(
        slide,
        image_path,
        Inches(1.0),
        Inches(1.25),
        Inches(11.3),
        Inches(5.7),
        caption,
    )


def slide_dual_image(
    prs: Presentation,
    title: str,
    left_path: Path,
    left_caption: str,
    right_path: Path,
    right_caption: str,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header_bar(slide, title)
    add_picture(slide, left_path, Inches(0.65), CONTENT_TOP, Inches(6.0), Inches(5.3), left_caption)
    add_picture(slide, right_path, Inches(6.85), CONTENT_TOP, Inches(6.0), Inches(5.3), right_caption)


def slide_text_only(prs: Presentation, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide)
    add_header_bar(slide, title)
    add_bullets(slide, bullets, width=Inches(11.8), font_size=22)


def slide_thanks(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, PRIMARY_DARK)
    box = slide.shapes.add_textbox(Inches(1.0), Inches(2.5), Inches(11.3), Inches(2.5))
    frame = box.text_frame
    frame.text = "敬请各位老师批评指正\n\n答辩人：张钧炀    导师：敖建锋"
    for index, paragraph in enumerate(frame.paragraphs):
        paragraph.font.size = Pt(32 if index == 0 else 20)
        paragraph.font.bold = index == 0
        paragraph.font.color.rgb = WHITE
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.space_after = Pt(16)


def build() -> None:
    presentation = Presentation()
    presentation.slide_width = SLIDE_W
    presentation.slide_height = SLIDE_H

    slide_cover(presentation)
    slide_contents(presentation)

    slide_section(presentation, "01", "绪论")
    slide_text_image(
        presentation,
        "研究背景与意义",
        [
            "激光雷达点云在测绘建模、地形量测中应用广泛",
            "离群点、悬浮点会干扰后续建模与精度评定",
            "SOR / ROR 是经典、可解释的预处理手段",
            "需要便于查看滤波效果的成果检查方式",
        ],
        img("image11.png"),
        "校园点云测区示例",
    )
    slide_text_only(
        presentation,
        "研究目标",
        [
            "实现 SOR、ROR 两种离群点滤波并比较参数影响",
            "从去除率、点数变化、结构保留等角度评价效果",
            "开发 Web 端程序，支持滤波前后三维对比查看",
            "扩展：分割、语义分类与语义精简（可选流程）",
        ],
    )

    slide_section(presentation, "02", "点云滤波理论基础")
    slide_text_image(
        presentation,
        "点云数据与 LAS 格式",
        [
            "点云由大量离散三维点组成，常含坐标与属性",
            "LAS 是机载/地面激光常用二进制格式",
            "处理前需建立空间索引以加速邻域搜索",
        ],
        img("image2.png"),
        "图2-1  LAS 文件结构",
    )
    slide_text_image(
        presentation,
        "邻域索引：KD-Tree 与 Ball Tree",
        [
            "KD-Tree：适合三维坐标邻域查询，用于 SOR",
            "Ball Tree：适合半径邻域统计，用于 ROR",
            "核心思想：先找邻域，再据统计或密度判离群",
        ],
        img("image3.png"),
        "图2-2  Ball tree 查询过程",
    )
    slide_text_image(
        presentation,
        "SOR 统计离群点滤波",
        [
            "计算每点 k 近邻平均距离",
            "若 mean_dist > μ + λ·σ，则判为离群点",
            "优点：抑制随机噪声；缺点：对 k、λ 敏感",
        ],
        img("image4.png"),
        "图3-1  统计离群滤波流程",
    )
    slide_text_image(
        presentation,
        "ROR 半径离群点滤波",
        [
            "统计半径 r 内邻域点数",
            "若邻域点数 < minN，则判为孤立离群点",
            "适合剔除稀疏悬浮点；对 r、minN 敏感",
        ],
        img("image5.png"),
        "图3-2  半径离群滤波流程（请核对）",
    )

    slide_section(presentation, "03", "程序设计与实现思路")
    slide_text_image(
        presentation,
        "程序总体思路（面向测绘成果）",
        [
            "输入：校园/城市尺度 LAS 点云",
            "处理：后台执行 SOR 或 ROR，输出 filtered.las",
            "查看：浏览器双窗对比原始与滤波结果",
            "创新：Web 端程序，免安装、易演示（区别于 Qt 桌面工具）",
        ],
        img("chapter3_overall_dataflow.png"),
        "系统总体数据流程（简述）",
    )
    slide_text_only(
        presentation,
        "主要创新点",
        [
            "① Web 端点云滤波与对比查看，降低使用与答辩演示门槛",
            "② SOR / ROR 参数化实验，给出去除率等量化指标",
            "③ 在滤波基础上扩展分割、语义分类与语义精简流程",
        ],
    )

    slide_section(presentation, "04", "实验结果与分析")
    slide_full_image(presentation, "原始点云", img("image11.png"), "图4-1  原始点云")
    slide_dual_image(
        presentation,
        "滤波效果对比",
        img("image13.png"),
        "图4-2  统计滤波（SOR）",
        img("image15.png"),
        "图4-3  半径滤波（ROR）",
    )
    slide_dual_image(
        presentation,
        "分割实验结果",
        img("image17.png"),
        "图4-5  密度聚类分割",
        img("image18.png"),
        "图4-6  区域生长分割",
    )
    slide_dual_image(
        presentation,
        "语义分类与精简",
        img("image20.png"),
        "图4-8  语义分类结果",
        img("image21.png"),
        "图4-9  语义精简结果",
    )
    slide_text_only(
        presentation,
        "参数敏感性（节选）",
        [
            "SOR · k：5→50，去除率约 10.09% → 8%（邻域越大越保守）",
            "SOR · std：0.5→2.0，去除率约 16.45% → 3.60%",
            "ROR · r：0.1→0.35，去除率约 37.35% → 0.52%",
            "ROR · minN：5→20，去除率约 1.73% → 9.29%",
            "详见论文表4-1～表4-4",
        ],
    )

    slide_section(presentation, "05", "结论与展望")
    slide_text_only(
        presentation,
        "结论与展望",
        [
            "结论：SOR/ROR 可有效去除离群点，主体几何结构基本保留",
            "结论：Web 端便于成果检查，适合毕设演示与分享",
            "不足：参数需结合场景调试；浏览器预览为抽稀显示",
            "展望：自适应参数、分块滤波、剖面与局部 ROI 对比",
        ],
    )
    slide_thanks(presentation)

    presentation.save(str(OUTPUT_PATH))
    print(f"Saved: {OUTPUT_PATH} ({len(presentation.slides)} slides)")


if __name__ == "__main__":
    build()
