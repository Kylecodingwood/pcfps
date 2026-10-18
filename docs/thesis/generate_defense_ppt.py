#!/usr/bin/env python3
"""Generate graduation defense PPT for PCFPS thesis."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

PRIMARY = RGBColor(0x1A, 0x56, 0x8E)
ACCENT = RGBColor(0x2E, 0x86, 0xC1)
TEXT = RGBColor(0x2C, 0x3E, 0x50)
LIGHT = RGBColor(0xF4, 0xF8, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_slide_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title_bar(slide, title: str) -> None:
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.95))
    bar.fill.solid()
    bar.fill.fore_color.rgb = PRIMARY
    bar.line.fill.background()
    text_frame = bar.text_frame
    text_frame.text = title
    paragraph = text_frame.paragraphs[0]
    paragraph.font.size = Pt(28)
    paragraph.font.bold = True
    paragraph.font.color.rgb = WHITE
    paragraph.alignment = PP_ALIGN.LEFT
    text_frame.margin_left = Inches(0.35)


def add_bullets(slide, items: list[str], left=0.55, top=1.25, width=8.9, height=5.8, font_size=20) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    text_frame = box.text_frame
    text_frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.size = Pt(font_size)
        paragraph.font.color.rgb = TEXT
        paragraph.space_after = Pt(10)


def add_two_column(slide, left_title: str, left_items: list[str], right_title: str, right_items: list[str]) -> None:
    left_box = slide.shapes.add_textbox(Inches(0.55), Inches(1.2), Inches(4.3), Inches(5.8))
    right_box = slide.shapes.add_textbox(Inches(5.15), Inches(1.2), Inches(4.3), Inches(5.8))
    for box, title, items in ((left_box, left_title, left_items), (right_box, right_title, right_items)):
        text_frame = box.text_frame
        title_paragraph = text_frame.paragraphs[0]
        title_paragraph.text = title
        title_paragraph.font.size = Pt(22)
        title_paragraph.font.bold = True
        title_paragraph.font.color.rgb = PRIMARY
        for item in items:
            paragraph = text_frame.add_paragraph()
            paragraph.text = item
            paragraph.font.size = Pt(18)
            paragraph.font.color.rgb = TEXT
            paragraph.space_after = Pt(8)


def build_presentation() -> Presentation:
    presentation = Presentation()
    presentation.slide_width = Inches(10)
    presentation.slide_height = Inches(7.5)

    slides_data: list[tuple[str, str, list[str] | None, dict | None]] = [
        ("cover", "", None, {"type": "cover"}),
        ("目录", "", [
            "1. 研究背景与意义",
            "2. 研究现状与目标",
            "3. 系统总体设计",
            "4. 核心算法与功能实现",
            "5. 实验结果与分析",
            "6. 结论与展望",
        ], None),
        ("研究背景与意义", "", [
            "激光雷达、三维扫描等技术推动点云数据在测绘、建模、自动驾驶等领域广泛应用",
            "点云预处理中，离群点与噪声点会干扰后续建模、量测与场景理解",
            "传统脚本/桌面工具（PCL、CloudCompare）操作分散，不利于任务管理与成果追溯",
            "本课题从测绘应用出发，设计并实现可任务化执行的点云滤波程序",
            "实现 SOR / ROR 滤波、Web 可视化对比及分割/分类等扩展能力",
        ], None),
        ("国内外研究现状", "", [
            "国外：统计滤波、中值滤波、曲面拟合等传统方法较早成熟；深度学习去噪鲁棒性更强",
            "国内：机载 LiDAR 地面滤波、城市复杂场景处理、点云重建优化研究进展迅速",
            "常见思路：基于邻域统计（SOR）与基于局部密度（ROR）的离群点剔除",
            "工程趋势：前后端分离、任务编排、浏览器轻量化预览成为实用化方向",
            "本课题定位：经典可解释算法 + 工程化平台，而非复刻通用点云工作站",
        ], None),
        ("研究目标与主要内容", "", [
            "目标：面向校园/城市尺度 LAS 点云，实现离群点去除与成果检查的一体化工具",
            "算法层：实现 SOR（cKDTree）与 ROR（BallTree）及参数化对比实验",
            "平台层：Java 任务管理 + Python 算法服务 + React/Three.js 双窗预览",
            "扩展层：欧式聚类/区域生长分割、PMF/CSF/规则语义分类、语义精简",
            "验证层：从去除率、点数变化、耗时、结构保真度及可视化效果多维度评估",
        ], None),
        ("系统总体架构", "", [
            "前后端分离 + 算法服务解耦的三层架构",
            "表现层：React + Ant Design + Three.js，任务向导与双窗点云工作台",
            "业务层：Spring Boot 3 + PostgreSQL，JWT 鉴权、任务调度、指标落库",
            "算法层：FastAPI 网关调度 python_filter / python_derived 子程序",
            "数据层：LAS 全精度产物落盘；浏览器仅加载后端抽稀后的 JSON 预览点",
            "设计原则：全精度在磁盘，Web 为可交互预览层，任务可追溯、可复现",
        ], None),
        ("开发环境与关键技术", "", None, {
            "type": "two_column",
            "left_title": "运行环境",
            "left_items": [
                "后端：Java 17、Spring Boot 3、PostgreSQL",
                "算法：Python 3、laspy、NumPy、SciPy、scikit-learn",
                "网关：FastAPI + 单 worker 线程池",
                "前端：Vite、React 18、TypeScript、Ant Design",
                "可视化：Three.js 点云渲染",
            ],
            "right_title": "关键技术",
            "right_items": [
                "KD-Tree / Ball Tree 加速邻域检索",
                "LAS 流式子集写出，保留 header 尺度",
                "任务状态机：PENDING → RUNNING → SUCCEEDED/FAILED",
                "comparison-bundle 聚合接口支撑双窗对比",
                "preview-points：解码 → 抽稀 → JSON",
            ],
        }),
        ("SOR 统计离群点滤波", "", [
            "思想：正常点与其 k 近邻距离稳定；离群点平均邻域距离显著偏大",
            "步骤：构建 cKDTree → 求每点 k 近邻均方根距离 → 计算全局均值与标准差",
            "判据：mean_distᵢ > μ + λ·σ 的点判为离群并剔除",
            "默认参数：k=20，std_ratio=1.0（可在任务中配置）",
            "特点：速度快、抑制随机噪声效果好；对 k、λ 及点密度变化较敏感",
            "实验：k 增大时去除率下降；std 增大时剔除更宽松",
        ], None),
        ("ROR 半径离群点滤波", "", [
            "思想：正常点在半径 r 内应有足够邻域点；孤立稀疏点邻域点数不足",
            "步骤：构建 BallTree → 半径查询统计邻域点数 → 与阈值 minN 比较",
            "判据：邻域点数（不含自身）< minN 则剔除",
            "默认参数：radius=1.0，min_neighbors=5",
            "特点：适合剔除孤立噪声；对 r、minN 及局部密度敏感",
            "实验：r 越大剔除越宽松；minN 增大时去除率上升但可能误删边缘点",
        ], None),
        ("系统功能与任务流程", "", [
            "用户登录 → 填写服务器 LAS 绝对路径 → 选择 SOR/ROR 及参数",
            "Java 创建任务并持久化参数快照 → 调用 Python 网关异步执行",
            "Python 写出 filtered.las，回传 N_before、N_after、去除比例、耗时",
            "前端轮询任务状态，成功后进入双窗对比页查看滤波前后效果",
            "派生任务：在滤波成功结果上继续分割、分类、语义精简",
            "支持任务列表、指标展示、产物下载与路径预览",
        ], None),
        ("扩展功能：分割、分类与精简", "", [
            "分割：欧式聚类（DBSCAN）、区域生长，输出带簇标签的 derived.las",
            "分类：PMF-lite、CSF-lite、校园规则伪语义四分类（地面/建筑/植被/其它）",
            "语义精简：保留地面体素抽稀与建筑完整点，去除植被/车辆等时相地物",
            "预览：classification 伪彩色模式便于观察簇/类别分布",
            "价值：由“干净几何点云”升级为“带语义、可工程简化”的成果",
        ], None),
        ("Web 端可视化展示", "", [
            "点云工作台：Segmented 切换滤波对比 / 派生双窗 / 路径双窗",
            "双窗联动：原始 vs 滤波、滤波 vs 分割/分类，同视角旋转缩放",
            "量化面板：全量点数、去除比例、Python 处理耗时（runMetrics）",
            "抽稀说明：页面标注 truncated 与采样规模，避免误解为全量显示",
            "答辩演示策略：全局看指标，局部看几何差异",
        ], None),
        ("滤波实验结果", "", [
            "数据：校园倾斜摄影点云 LAS，包含建筑、道路、植被等区域",
            "SOR：k=5 时去除率约 10.09%；k=50 时降至约 8%，邻域越大越保守",
            "SOR：std 系数增大时剔除阈值放宽，结构保留更好",
            "ROR：r=0.1 时去除率可达 37.35%；r=0.35 时仅 0.52%",
            "ROR：minN 由 5 增至 20，去除率由 1.73% 升至 9.29%",
            "对比：SOR/ROR 均能有效去除边缘悬浮点，主体几何结构基本保持",
        ], None),
        ("分割与分类实验结果", "", [
            "分割：欧式聚类可将建筑条带、屋顶等连续区域划分为独立簇",
            "分割：区域生长适用于法向相近的局部连通区域提取",
            "分类：校园约 8000 万点场景，规则语义分类后可按类别统计与导出",
            "精简：去除植被/车辆后，建筑细节保留，场景更利于 BIM/建模对接",
            "对比：语义结果与 CloudCompare 人工框选基本一致",
            "局限：浏览器均匀抽稀下，小尺度细节无法完整呈现",
        ], None),
        ("主要创新点", "", [
            "① 经典 SOR/ROR 与空间索引结合，形成参数化、任务化、可追溯滤波流程",
            "② 在滤波链路上扩展 RF 语义分类与语义裁剪，支撑工程化成果输出",
            "③ Web 任务管理 + 双窗预览，降低桌面工具操作门槛，便于答辩演示",
            "④ 全精度 LAS 与 JSON 预览分层，兼顾精度保留与浏览器可交互性",
            "⑤ 滤波—分割—分类—精简一体化，覆盖测绘点云预处理常见链路",
        ], None),
        ("结论", "", [
            "完成了面向测绘场景的点云滤波程序设计与实现，系统运行稳定",
            "SOR/ROR 可有效去除零星孤立噪声，并输出量化指标便于参数对比",
            "任务调度、Python 服务调用、产物登记与预览展示链路完整闭环",
            "分割、语义分类与精简功能验证了从几何去噪到语义成果扩展的可行性",
            "达到了毕业设计预期目标，具备一定的实用与展示价值",
        ], None),
        ("不足与展望", "", [
            "不足：复杂场景边缘/植被粘连仍有残留或误删；语义依赖规则与样本分布",
            "不足：超大 LAS 整文件加载压力大；评测以视觉与基础指标为主",
            "展望：多尺度邻域、自适应阈值、分块/瓦片化并行滤波",
            "展望：剖面切片、ROI 局部对比、视锥加密抽稀等预览增强",
            "展望：与 BIM/实景三维生产流程 deeper 对接，服务智慧城市与数字孪生",
        ], None),
        ("致谢", "", [
            "感谢各位老师在毕业设计过程中的悉心指导！",
            "感谢同学与朋友在测试与演示中的帮助！",
            "恳请各位老师批评指正！",
            "",
            "测绘223  张钧炀",
            f"{date.today().year}年{date.today().month}月",
        ], {"type": "thanks"}),
    ]

    for title, subtitle, bullets, extra in slides_data:
        slide_layout = presentation.slide_layouts[6]
        slide = presentation.slides.add_slide(slide_layout)
        set_slide_bg(slide, LIGHT)

        if extra and extra.get("type") == "cover":
            set_slide_bg(slide, PRIMARY)
            title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(8.4), Inches(1.2))
            title_frame = title_box.text_frame
            title_frame.text = "点云数据滤波程序设计"
            title_paragraph = title_frame.paragraphs[0]
            title_paragraph.font.size = Pt(40)
            title_paragraph.font.bold = True
            title_paragraph.font.color.rgb = WHITE
            title_paragraph.alignment = PP_ALIGN.CENTER

            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(3.0), Inches(8.4), Inches(0.8))
            sub_frame = sub_box.text_frame
            sub_frame.text = "毕业设计答辩"
            sub_paragraph = sub_frame.paragraphs[0]
            sub_paragraph.font.size = Pt(24)
            sub_paragraph.font.color.rgb = RGBColor(0xD6, 0xEA, 0xF8)
            sub_paragraph.alignment = PP_ALIGN.CENTER

            info_box = slide.shapes.add_textbox(Inches(0.8), Inches(4.5), Inches(8.4), Inches(1.8))
            info_frame = info_box.text_frame
            info_frame.text = "答辩人：测绘223  张钧炀\n指导教师：（请填写）\n学校/学院：（请填写）"
            for paragraph in info_frame.paragraphs:
                paragraph.font.size = Pt(20)
                paragraph.font.color.rgb = WHITE
                paragraph.alignment = PP_ALIGN.CENTER
                paragraph.space_after = Pt(6)
            continue

        if extra and extra.get("type") == "thanks":
            set_slide_bg(slide, PRIMARY)
            box = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(8.0), Inches(3.5))
            text_frame = box.text_frame
            text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            for index, item in enumerate(bullets or []):
                paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
                paragraph.text = item
                paragraph.font.size = Pt(28 if index == 0 else 22)
                paragraph.font.bold = index == 0
                paragraph.font.color.rgb = WHITE
                paragraph.alignment = PP_ALIGN.CENTER
                paragraph.space_after = Pt(12)
            continue

        add_title_bar(slide, title)
        if extra and extra.get("type") == "two_column":
            add_two_column(
                slide,
                extra["left_title"],
                extra["left_items"],
                extra["right_title"],
                extra["right_items"],
            )
        elif bullets:
            add_bullets(slide, bullets)

    return presentation


def main() -> None:
    output_dir = Path("/Users/kyle/Desktop")
    output_path = output_dir / "毕业设计答辩PPT-测绘223张钧炀.pptx"
    presentation = build_presentation()
    presentation.save(output_path)
    print(f"Generated: {output_path}")


if __name__ == "__main__":
    main()
