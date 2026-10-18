#!/usr/bin/env python3
"""Fill 江理答辩5 template — text pages keep décor; image pages get clean layout."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

TEMPLATE_PATH = Path("/Users/kyle/Downloads/江理答辩5.pptx")
OUTPUT_PATH = Path("/Users/kyle/Desktop/毕业设计答辩PPT-测绘223张钧炀.pptx")
REPO_IMAGES = Path(__file__).resolve().parent / "images"
DOCX_IMAGES = Path(__file__).resolve().parent / "extracted_docx_media"

OLD_RED = RGBColor(0xD0, 0x40, 0x27)
PRIMARY = RGBColor(0x3D, 0x6B, 0x8A)
TEXT_DARK = RGBColor(0x33, 0x33, 0x33)
MUTED = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0xB8, 0x5A, 0x4A)

CONTENT_TOP = Inches(1.25)
HEADER_BOTTOM_IN = 1.15

PLACEHOLDER_PREFIXES = (
    "点击",
    "在此",
    "请在此",
    "请在这里",
    "详写",
    "添加",
    "输入",
    "单击",
    "您的内容",
    "标题文字",
)
PLACEHOLDER_EXACT = {
    "绪论",
    "研究方法与研究思路",
    "关键技术与实践难点",
    "研究成果与应用",
    "相关建议及总结",
}


def img(name: str) -> Path:
    candidate = DOCX_IMAGES / name
    return candidate if candidate.exists() else REPO_IMAGES / name


def rgb_equal(first: RGBColor, second: RGBColor) -> bool:
    return int(first) == int(second)


def soften_slide_colors(slide) -> None:
    for shape in slide.shapes:
        try:
            if shape.fill.type is not None and shape.fill.fore_color.type is not None:
                if rgb_equal(shape.fill.fore_color.rgb, OLD_RED):
                    shape.fill.fore_color.rgb = PRIMARY
        except Exception:
            pass
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for child in shape.shapes:
                try:
                    if child.fill.type is not None and child.fill.fore_color.type is not None:
                        if rgb_equal(child.fill.fore_color.rgb, OLD_RED):
                            child.fill.fore_color.rgb = PRIMARY
                except Exception:
                    pass


def is_placeholder_text(text: str) -> bool:
    stripped = text.strip()
    if not stripped:
        return False
    if stripped in PLACEHOLDER_EXACT:
        return True
    return any(stripped.startswith(prefix) for prefix in PLACEHOLDER_PREFIXES)


def set_shape_text(
    shape,
    text: str,
    size: int = 16,
    bold: bool = False,
    color: RGBColor = TEXT_DARK,
) -> None:
    if not hasattr(shape, "text_frame"):
        return
    shape.text = text
    for paragraph in shape.text_frame.paragraphs:
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color


def iter_text_shapes(slide):
    for shape in slide.shapes:
        if hasattr(shape, "text_frame"):
            yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for child in shape.shapes:
                if hasattr(child, "text_frame"):
                    yield child


def collect_placeholder_shapes(slide) -> list:
    return [shape for shape in iter_text_shapes(slide) if is_placeholder_text(shape.text)]


def fill_placeholders_ordered(slide, values: list[str], size: int = 16) -> None:
    for shape, value in zip(collect_placeholder_shapes(slide), values):
        set_shape_text(shape, value, size=size)


def replace_exact_text(slide, mapping: dict[str, str], size: int = 18) -> None:
    for shape in iter_text_shapes(slide):
        original = shape.text.strip()
        for key, value in mapping.items():
            if original == key or key in original:
                set_shape_text(shape, value, size=size)
                break


def set_section_title(slide, section_label: str, intro: str = "") -> None:
    replaced = False
    for shape in iter_text_shapes(slide):
        text = shape.text.strip()
        if text in {"01", "02", "03", "04", "05"}:
            continue
        if not replaced and any(
            keyword in text
            for keyword in (
                "绪论",
                "研究方法",
                "关键技术",
                "研究成果",
                "相关建议",
                "点云滤波",
                "程序设计",
                "实验结果",
                "结论",
            )
        ):
            set_shape_text(shape, section_label, size=24, bold=True, color=WHITE)
            replaced = True
    for shape in collect_placeholder_shapes(slide):
        set_shape_text(shape, intro if intro else section_label, size=14)


def fill_toc_groups(slide, chapter_titles: list[str]) -> None:
    group_shapes = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.GROUP]
    for group_shape, title in zip(group_shapes[5:10], chapter_titles):
        for child in group_shape.shapes:
            if hasattr(child, "text_frame") and child.text.strip():
                set_shape_text(child, title, size=18, bold=True)
                break


def is_header_decoration(shape) -> bool:
    top_inches = shape.top.inches
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and top_inches < 0.35:
        return True
    if shape.shape_type == MSO_SHAPE_TYPE.LINE:
        return True
    if top_inches < 0.35 and shape.width.inches < 0.55:
        return True
    if top_inches < HEADER_BOTTOM_IN and shape.width.inches > 2.5 and shape.height.inches < 1.25:
        return True
    if shape.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER and top_inches < HEADER_BOTTOM_IN:
        return True
    return False


def clear_content_shapes(slide) -> None:
    """Remove body décor so image pages have a clean canvas below the header."""
    for shape in list(slide.shapes):
        if is_header_decoration(shape):
            continue
        shape.element.getparent().remove(shape.element)


def add_page_title(slide, title: str) -> None:
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER and shape.top.inches < HEADER_BOTTOM_IN:
            set_shape_text(shape, title, size=22, bold=True, color=PRIMARY)
            return
    title_box = slide.shapes.add_textbox(Inches(0.9), Inches(0.38), Inches(10), Inches(0.5))
    set_shape_text(title_box, title, size=22, bold=True, color=PRIMARY)


def add_bullets(slide, items: list[str], left, top, width, height, size: int = 18) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = TEXT_DARK
        paragraph.space_after = Pt(8)


def place_picture(slide, image_path: Path, left, top, width, height, caption: str = "") -> None:
    if image_path.exists():
        slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)
    else:
        placeholder = slide.shapes.add_shape(1, left, top, width, height)
        placeholder.fill.solid()
        placeholder.fill.fore_color.rgb = RGBColor(0xE8, 0xED, 0xF2)
        placeholder.line.color.rgb = PRIMARY
        text_frame = placeholder.text_frame
        text_frame.text = f"【待补图】\n{caption}"
        text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    if caption:
        cap_box = slide.shapes.add_textbox(left, top + height + Inches(0.04), width, Inches(0.28))
        cap_paragraph = cap_box.text_frame.paragraphs[0]
        cap_paragraph.text = caption
        cap_paragraph.font.size = Pt(11)
        cap_paragraph.font.color.rgb = MUTED
        cap_paragraph.alignment = PP_ALIGN.CENTER


def layout_full_image(slide, title: str, image_path: Path, caption: str) -> None:
    clear_content_shapes(slide)
    add_page_title(slide, title)
    place_picture(slide, image_path, Inches(1.0), CONTENT_TOP, Inches(11.3), Inches(5.55), caption)


def layout_dual_image(
    slide,
    title: str,
    left_path: Path,
    left_caption: str,
    right_path: Path,
    right_caption: str,
) -> None:
    clear_content_shapes(slide)
    add_page_title(slide, title)
    place_picture(slide, left_path, Inches(0.65), CONTENT_TOP, Inches(5.95), Inches(5.2), left_caption)
    place_picture(slide, right_path, Inches(6.75), CONTENT_TOP, Inches(5.95), Inches(5.2), right_caption)


def layout_text_image(
    slide,
    title: str,
    bullets: list[str],
    image_path: Path,
    caption: str,
) -> None:
    clear_content_shapes(slide)
    add_page_title(slide, title)
    add_bullets(slide, bullets, Inches(0.65), CONTENT_TOP, Inches(5.5), Inches(5.5), size=17)
    place_picture(slide, image_path, Inches(6.55), CONTENT_TOP, Inches(6.15), Inches(5.2), caption)


def layout_quad_image(
    slide,
    title: str,
    top_left: tuple[Path, str],
    top_right: tuple[Path, str],
    bottom_left: tuple[Path, str],
    bottom_right: tuple[Path, str],
) -> None:
    clear_content_shapes(slide)
    add_page_title(slide, title)
    row_top = CONTENT_TOP
    row_bottom = Inches(4.15)
    col_left = Inches(0.65)
    col_right = Inches(6.75)
    image_width = Inches(5.95)
    image_height = Inches(2.35)
    for column, row, (path, caption) in (
        (col_left, row_top, top_left),
        (col_right, row_top, top_right),
        (col_left, row_bottom, bottom_left),
        (col_right, row_bottom, bottom_right),
    ):
        place_picture(slide, path, column, row, image_width, image_height, caption)


def build() -> None:
    presentation = Presentation(str(TEMPLATE_PATH))
    for slide in presentation.slides:
        soften_slide_colors(slide)

    # ── 模板页：只改文字，不加图 ─────────────────────────────
    cover = presentation.slides[0]
    replace_exact_text(
        cover,
        {
            "毕业论文答辩PPT": "点云数据滤波程序设计",
            "XX学院XX系XX级XX班": "低空技术学院  测绘工程223班",
            "答辩人：  导师：": "答辩人：张钧炀    导师：敖建锋",
        },
        size=20,
    )
    for shape in iter_text_shapes(cover):
        if "点云数据滤波程序设计" in shape.text:
            set_shape_text(shape, "点云数据滤波程序设计", size=32, bold=True, color=WHITE)

    fill_toc_groups(
        presentation.slides[1],
        [
            "研究背景与意义",
            "点云滤波理论基础",
            "程序设计与实现思路",
            "实验结果与分析",
            "结论与展望",
        ],
    )

    set_section_title(presentation.slides[2], "01  绪论", "激光点云去噪是测绘建模预处理的重要环节")

    fill_placeholders_ordered(
        presentation.slides[3],
        [
            "激光点云应用",
            "机载/地面激光在测绘建模、地形量测中应用广泛",
            "噪声与离群点",
            "随机噪声、悬浮点影响 DEM/三维建模精度",
            "经典滤波方法",
            "SOR / ROR 可解释、参数可控，适合工程预处理",
            "成果查看需求",
            "需要直观对比滤波前后点云，便于检查与答辩演示",
            "研究背景",
            "激光点云应用",
            "噪声与离群点",
            "经典滤波方法",
        ],
        size=14,
    )

    goals_slide = presentation.slides[4]
    replace_exact_text(goals_slide, {"添加标题": "研究目标"}, size=20)
    fill_placeholders_ordered(
        goals_slide,
        [
            "实现 SOR、ROR 两种离群点滤波并比较参数影响",
            "从去除率、点数变化、结构保留等角度评价效果",
            "开发 Web 端程序，支持滤波前后三维对比查看",
            "扩展：分割、语义分类与语义精简（可选）",
            "形成可复现的校园点云滤波实验与成果展示",
        ],
        size=15,
    )

    set_section_title(presentation.slides[5], "02  点云滤波理论基础")

    fill_placeholders_ordered(
        presentation.slides[7],
        [
            "邻域索引",
            "KD-Tree 适合 k 近邻查询（SOR）",
            "Ball Tree",
            "适合半径邻域统计（ROR）",
            "滤波思想",
            "先建立索引，再按统计或密度判断离群点",
            "邻域索引与滤波思想",
            "邻域索引",
            "Ball Tree",
            "滤波思想",
        ],
        size=14,
    )

    fill_placeholders_ordered(
        presentation.slides[8],
        [
            "SOR",
            "k 近邻距离统计，剔除偏离分布的点",
            "ROR",
            "半径内邻域点数不足则判为离群点",
            "SOR 统计离群滤波",
            "mean_dist > μ + λσ 判为离群点",
            "ROR 半径离群滤波",
            "邻域点数 < minN 判为孤立离群点",
        ],
        size=14,
    )

    set_section_title(
        presentation.slides[9],
        "03  程序设计与实现思路",
        "面向测绘成果：输入 LAS → 滤波 → 对比查看",
    )

    fill_placeholders_ordered(
        presentation.slides[10],
        [
            "选择数据",
            "上传/选择校园 LAS 与滤波参数",
            "执行滤波",
            "后台 SOR 或 ROR 处理",
            "输出成果",
            "生成 filtered.las 等成果文件",
            "对比查看",
            "浏览器双窗三维对比（Web 端创新点）",
            "输入文本",
            "Web 端易演示；总体流程：LAS → 滤波 → 双窗对比",
        ],
        size=13,
    )

    fill_placeholders_ordered(
        presentation.slides[12],
        [
            "密度聚类分割",
            "在滤波结果上按密度分离地物",
            "区域生长分割",
            "由种子点扩展连通区域",
            "语义分类",
            "对分割块赋予类别标签",
            "语义精简",
            "按类别裁剪保留目标地物",
            "扩展流程",
            "滤波 → 分割 → 分类 → 精简",
            "面向校园点云成果检查",
        ],
        size=13,
    )

    set_section_title(presentation.slides[13], "04  实验结果与分析")

    innovation_slide = presentation.slides[15]
    fill_placeholders_ordered(
        innovation_slide,
        [
            "易演示",
            "Web 端",
            "可量化",
            "参数实验",
            "可扩展",
            "处理链",
            "Web 端点云滤波与对比查看，区别于 Qt 桌面工具",
            "创新一",
            "SOR / ROR 参数化实验与去除率等量化指标",
            "创新二",
            "滤波—分割—语义处理一体化扩展流程",
            "创新三",
        ],
        size=12,
    )
    for shape in iter_text_shapes(innovation_slide):
        if shape.top.inches > 4.3 and shape.left.inches > 5.0 and shape.left.inches < 6.0:
            set_shape_text(shape, "主要创新点", size=18, bold=True, color=PRIMARY)

    set_section_title(presentation.slides[17], "05  结论与展望")

    fill_placeholders_ordered(
        presentation.slides[19],
        [
            "SOR · k：5→50，去除率约 10.09% → 8%",
            "参数实验",
            "SOR · std：0.5→2.0，去除率约 16.45% → 3.60%",
            "参数实验",
            "ROR · r：0.1→0.35，去除率约 37.35% → 0.52%",
            "参数实验",
            "ROR · minN：5→20，去除率约 1.73% → 9.29%",
            "参数实验",
            "SOR/ROR 可有效去除离群点，结构保留较好",
            "结论",
            "Web 端便于滤波前后对比与答辩演示",
            "创新",
            "自适应参数、分块滤波、剖面与局部对比",
            "展望",
            "面向测绘生产的预处理工具链持续完善",
            "展望",
        ],
        size=11,
    )

    thanks_slide = presentation.slides[20]
    replace_exact_text(
        thanks_slide,
        {
            "铭记师恩  感谢有您": "敬请各位老师批评指正",
            "XX学院XX系XX级XX班": "低空技术学院  测绘工程223班",
            "答辩人：  导师：": "答辩人：张钧炀    导师：敖建锋",
        },
        size=18,
    )
    for shape in iter_text_shapes(thanks_slide):
        if "敬请各位老师批评指正" in shape.text:
            set_shape_text(shape, "敬请各位老师批评指正", size=30, bold=True, color=WHITE)

    # ── 图片页：清空内容区后自主排版（模板装饰与图不混排）──────────
    # slide 7  阶梯装饰 → 整页 LAS 图
    layout_full_image(
        presentation.slides[6],
        "点云数据与 LAS 格式",
        img("image2.png"),
        "图2-1  LAS 文件结构",
    )

    # slide 8  分组模板 → 仅文字（Ball tree 概念），不放图
    # slide 9  圆环模板 → 仅文字（SOR/ROR 原理），不放图

    # slide 12 六格装饰 → 双图（SOR/ROR 流程）；程序思路见 slide 11 时间轴文字
    layout_dual_image(
        presentation.slides[11],
        "SOR 与 ROR 滤波流程",
        img("image4.png"),
        "图3-1  SOR 统计离群滤波",
        img("image5.png"),
        "图3-2  ROR 半径离群滤波（请核对）",
    )

    # slide 13 六栏 → 仅文字（扩展功能），不放图

    # slide 15 圆环装饰 → 双图对比（原始 vs SOR）
    layout_dual_image(
        presentation.slides[14],
        "原始点云与 SOR 滤波效果",
        img("image11.png"),
        "图4-1  原始点云",
        img("image13.png"),
        "图4-2  统计滤波（SOR）",
    )

    # slide 16 创新环绕 → 仅文字，已在上方填充

    # slide 17 四象限 → 2×2 实验图（分割 + 语义）
    layout_quad_image(
        presentation.slides[16],
        "分割与语义实验结果",
        (img("image17.png"), "图4-5  密度聚类分割"),
        (img("image18.png"), "图4-6  区域生长分割"),
        (img("image20.png"), "图4-8  语义分类结果"),
        (img("image21.png"), "图4-9  语义精简结果"),
    )

    # slide 19 双栏模板 → 双图对比（滤波实验效果）
    layout_dual_image(
        presentation.slides[18],
        "SOR 与 ROR 滤波效果对比",
        img("image13.png"),
        "图4-2  统计滤波（SOR）",
        img("image15.png"),
        "图4-3  半径滤波（ROR）",
    )

    # slide 20 六宫格 → 仅文字（参数实验 + 结论），不放图

    presentation.save(str(OUTPUT_PATH))
    print(f"Saved: {OUTPUT_PATH} ({len(presentation.slides)} slides)")


if __name__ == "__main__":
    build()
